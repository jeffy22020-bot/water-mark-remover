from app.core import office


def test_pptx_watermarks_removed(samples, tmp_path):
    src = samples / "slides-evaluation.pptx"
    analysis = office.analyze(str(src))
    selected = [c for c in analysis["candidates"] if c["selected"]]
    assert any("Evaluation Warning" in c["label"] for c in selected)
    assert analysis["media"], "embedded pictures are listed"
    out = tmp_path / "clean.pptx"
    office.remove(str(src), str(out), {analysis["_keys"][c["id"]] for c in selected})

    from pptx import Presentation

    prs = Presentation(str(out))
    texts = [s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame]
    assert not any("Evaluation Warning" in t for t in texts)
    assert any("Project update" in t for t in texts)


def test_docx_header_watermark_removed(samples, tmp_path):
    src = samples / "agreement-draft.docx"
    analysis = office.analyze(str(src))
    assert analysis["candidates"][0]["label"] == "Watermark “DRAFT”"
    out = tmp_path / "clean.docx"
    office.remove(str(src), str(out), {analysis["_keys"][c["id"]] for c in analysis["candidates"]})
    assert office.analyze(str(out))["candidates"] == []

    from docx import Document

    doc = Document(str(out))
    assert doc.sections[0].header.paragraphs[0].text == "Acme Corp."


def test_replace_media(samples, tmp_path):
    import io

    from PIL import Image

    src = samples / "slides-evaluation.pptx"
    name = office.analyze(str(src))["media"][0]["name"]
    buf = io.BytesIO()
    Image.new("RGB", (260, 80), (0, 200, 0)).save(buf, "PNG")
    out = tmp_path / "replaced.pptx"
    office.replace_media(str(src), str(out), name, buf.getvalue())
    blob = office.get_media(str(out), name)
    assert Image.open(io.BytesIO(blob)).convert("RGB").getpixel((5, 5)) == (0, 200, 0)
