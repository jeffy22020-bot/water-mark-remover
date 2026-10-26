"""Generate the watermarked sample files shown under "Try one of these".

The same files are used as test fixtures, so every sample also returns the ground-truth
watermark mask where that makes sense. Everything is drawn procedurally: no third-party
images are bundled.

    python scripts/make_samples.py            # writes ./samples
"""

from __future__ import annotations

import sys
import zlib
from pathlib import Path

import cv2
import numpy as np

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.core.ffmpeg import FrameWriter, ffmpeg_exe, popen_kwargs  # noqa: E402


# ------------------------------------------------------------------ scenery
def _noise1d(n: int, octaves: int, rng: np.random.Generator) -> np.ndarray:
    out = np.zeros(n, np.float32)
    for o in range(octaves):
        pts = 2 ** (o + 2) + 1
        ctrl = rng.uniform(-1, 1, pts).astype(np.float32)
        out += np.interp(np.linspace(0, pts - 1, n), np.arange(pts), ctrl) / (2 ** o)
    return out


def landscape(width: int, height: int, seed: int = 3, sun=(0.72, 0.38)) -> np.ndarray:
    rng = np.random.default_rng(seed)
    y = np.linspace(0, 1, height, dtype=np.float32)[:, None]
    top = np.array([120, 60, 40], np.float32)  # BGR: deep blue-violet
    mid = np.array([150, 120, 240], np.float32)  # pink
    low = np.array([80, 190, 255], np.float32)  # orange
    t = np.clip(y / 0.62, 0, 1)
    sky = np.where(t[..., None] < 0.6, top + (mid - top) * (t[..., None] / 0.6), mid + (low - mid) * ((t[..., None] - 0.6) / 0.4))
    img = np.repeat(sky, width, axis=1).astype(np.float32)
    # sun glow
    yy, xx = np.mgrid[0:height, 0:width].astype(np.float32)
    d = np.sqrt((xx - sun[0] * width) ** 2 + (yy - sun[1] * height) ** 2) / max(width, height)
    img += (np.exp(-(d / 0.05) ** 2) * 120 + np.exp(-(d / 0.25) ** 2) * 50)[..., None] * np.array([0.6, 0.9, 1.0], np.float32)
    # clouds
    clouds = cv2.GaussianBlur(rng.uniform(0, 1, (height // 8, width // 8)).astype(np.float32), (0, 0), 2)
    clouds = cv2.resize(clouds, (width, height), interpolation=cv2.INTER_CUBIC)
    img += (np.clip(clouds - 0.55, 0, 1) * 160 * (1 - y))[..., None]
    # mountain layers
    colors = [(110, 80, 120), (90, 60, 90), (60, 45, 60)]
    for i, color in enumerate(colors):
        ridge = (0.45 + 0.07 * i + 0.09 * _noise1d(width, 6, rng)) * height
        mask = yy > ridge[None, :]
        shade = np.array(color, np.float32) * (0.9 + 0.1 * rng.uniform(0, 1, (height, width))[..., None].astype(np.float32))
        img = np.where(mask[..., None], shade, img)
    # lake with reflection and ripples
    horizon = int(height * 0.7)
    water = img[horizon - (height - horizon):horizon][::-1].copy() * 0.75 + np.array([60, 30, 10], np.float32)
    ripple = (np.sin(yy[horizon:] * 0.9 + np.sin(xx[horizon:] * 0.05) * 3) * 6)[..., None]
    img[horizon:] = water[: height - horizon] + ripple
    grain = rng.normal(0, 4, img.shape).astype(np.float32)
    return np.clip(img + grain, 0, 255).astype(np.uint8)


def _text_layer(shape, text, scale, thickness, positions, angle=0.0):
    layer = np.zeros(shape[:2], np.uint8)
    font = cv2.FONT_HERSHEY_DUPLEX
    for x, y in positions:
        cv2.putText(layer, text, (int(x), int(y)), font, scale, 255, thickness, cv2.LINE_AA)
    if angle:
        h, w = shape[:2]
        rot = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
        layer = cv2.warpAffine(layer, rot, (w, h))
    return layer


def _blend(img, layer, color=(255, 255, 255), opacity=0.55):
    a = (layer.astype(np.float32) / 255.0 * opacity)[..., None]
    out = img.astype(np.float32) * (1 - a) + np.array(color, np.float32) * a
    return np.clip(out, 0, 255).astype(np.uint8)


def sample_tiled_image(width=1280, height=853):
    clean = landscape(width, height, seed=3)
    positions = [(x, y) for y in range(-200, height + 400, 170) for x in range(-300, width + 300, 420)]
    layer = _text_layer(clean.shape, "(c) SAMPLE", 1.3, 3, positions, angle=25)
    return clean, _blend(clean, layer, opacity=0.5), layer


def sample_logo_image(width=1200, height=800):
    clean = landscape(width, height, seed=11, sun=(0.3, 0.42))
    layer = np.zeros(clean.shape[:2], np.uint8)
    cv2.circle(layer, (width - 250, height - 62), 22, 255, 5, cv2.LINE_AA)
    cv2.putText(layer, "PhotoStock", (width - 215, height - 48), cv2.FONT_HERSHEY_DUPLEX, 1.2, 255, 2, cv2.LINE_AA)
    return clean, _blend(clean, layer, opacity=0.85), layer


# ------------------------------------------------------------------ video
def write_sample_video(path: Path, seconds: float = 4.0, fps: int = 25, width: int = 640, height: int = 360, moving: bool = True):
    """Panning scene with a watermark that jumps corners half-way (TikTok style).

    Returns ``[(start, end, mask), ...]`` ground truth.
    """
    pano = landscape(width * 2, height, seed=5)
    frames = int(seconds * fps)
    logo_a = np.zeros((height, width), np.uint8)
    cv2.putText(logo_a, "@clipmaker", (14, 34), cv2.FONT_HERSHEY_DUPLEX, 0.9, 255, 2, cv2.LINE_AA)
    logo_b = np.zeros((height, width), np.uint8)
    cv2.putText(logo_b, "@clipmaker", (width - 200, height - 20), cv2.FONT_HERSHEY_DUPLEX, 0.9, 255, 2, cv2.LINE_AA)
    half = seconds / 2 if moving else seconds
    audio = path.with_suffix(".wav")
    _write_tone(audio, seconds)
    writer = FrameWriter(str(path), width, height, fps, audio_source=str(audio), crf=20)
    for i in range(frames):
        t = i / fps
        x = int((width) * i / frames)
        frame = pano[:, x:x + width].copy()
        # a boat crossing so the background is clearly changing
        bx = int(width * (0.1 + 0.8 * i / frames))
        cv2.ellipse(frame, (bx, int(height * 0.78)), (40, 10), 0, 0, 180, (30, 30, 40), -1)
        layer = logo_a if t < half else logo_b
        writer.write(_blend(frame, layer, opacity=0.8))
    writer.close()
    audio.unlink(missing_ok=True)
    if moving:
        return [(0.0, half, logo_a), (half, seconds, logo_b)]
    return [(0.0, seconds, logo_a)]


def _write_tone(path: Path, seconds: float, rate: int = 22050):
    import wave

    t = np.arange(int(seconds * rate)) / rate
    data = (np.sin(2 * np.pi * 440 * t) * 0.2 * 32767).astype(np.int16)
    with wave.open(str(path), "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(data.tobytes())


# ------------------------------------------------------------------ PDF
BODY = (
    "Quarterly report. Revenue grew steadily across all regions while costs stayed flat. "
    "The team shipped three major releases and onboarded forty new customers this quarter."
)


def write_sample_pdf(path: Path, pages: int = 3) -> None:
    """A PDF with three kinds of watermark: rotated transparent text, a repeated logo
    image and an /Artifact /Watermark block (the way Acrobat stamps them)."""
    import pikepdf
    from pikepdf import Array, Dictionary, Name

    pdf = pikepdf.new()
    font = pdf.make_indirect(Dictionary(Type=Name.Font, Subtype=Name.Type1, BaseFont=Name.Helvetica, Encoding=Name.WinAnsiEncoding))
    bold = pdf.make_indirect(Dictionary(Type=Name.Font, Subtype=Name("/Type1"), BaseFont=Name("/Helvetica-Bold"), Encoding=Name.WinAnsiEncoding))
    gs = pdf.make_indirect(Dictionary(Type=Name.ExtGState, ca=0.25, CA=0.25))
    logo = np.zeros((40, 160, 3), np.uint8)
    logo[:] = (40, 40, 200)
    cv2.putText(logo, "ACME DRAFT", (6, 28), cv2.FONT_HERSHEY_DUPLEX, 0.75, (255, 255, 255), 1, cv2.LINE_AA)
    rgb = cv2.cvtColor(logo, cv2.COLOR_BGR2RGB)
    image = pikepdf.Stream(pdf, zlib.compress(rgb.tobytes()))
    image.Type, image.Subtype = Name.XObject, Name.Image
    image.Width, image.Height = 160, 40
    image.ColorSpace, image.BitsPerComponent, image.Filter = Name.DeviceRGB, 8, Name.FlateDecode
    image = pdf.make_indirect(image)

    for n in range(pages):
        lines = []
        y = 760
        for paragraph in range(6):
            words = BODY.split()
            line = ""
            for word in words:
                if len(line) + len(word) > 80:
                    lines.append(f"BT /F1 11 Tf 72 {y} Td ({line.strip()}) Tj ET")
                    y -= 15
                    line = ""
                line += word + " "
            lines.append(f"BT /F1 11 Tf 72 {y} Td ({line.strip()}) Tj ET")
            y -= 28
        heading = f"BT /F2 18 Tf 72 790 Td (Page {n + 1} - Company Report) Tj ET"
        watermark_text = "q /GS1 gs 0.6 0.6 0.6 rg BT /F2 72 Tf 0.7071 0.7071 -0.7071 0.7071 150 250 Tm (CONFIDENTIAL) Tj ET Q"
        logo_draw = "q 160 0 0 40 410 30 cm /Im1 Do Q"
        artifact = "/Artifact <</Type /Pagination /Subtype /Watermark>> BDC q 0.9 0.2 0.2 rg BT /F2 14 Tf 230 60 Td (SAMPLE COPY - DO NOT DISTRIBUTE) Tj ET Q EMC"
        content = "\n".join([heading, *lines, watermark_text, logo_draw, artifact])
        page = pikepdf.Page(
            Dictionary(
                Type=Name.Page,
                MediaBox=Array([0, 0, 612, 842]),
                Resources=Dictionary(
                    Font=Dictionary(F1=font, F2=bold),
                    ExtGState=Dictionary(GS1=gs),
                    XObject=Dictionary(Im1=image),
                ),
                Contents=pdf.make_stream(content.encode("latin-1")),
            )
        )
        pdf.pages.append(page)
    pdf.save(path)


# ------------------------------------------------------------------ Office
def write_sample_pptx(path: Path, slides: int = 3) -> None:
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    tmp_logo = path.with_name("_logo.png")
    logo = np.zeros((80, 260, 4), np.uint8)
    cv2.putText(logo, "TRIAL LOGO", (8, 55), cv2.FONT_HERSHEY_DUPLEX, 1.3, (60, 60, 230, 255), 3, cv2.LINE_AA)
    cv2.imwrite(str(tmp_logo), logo)
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn

    # python-pptx can't add pictures to a master directly: add it to a scratch slide and move it.
    master = prs.slide_masters[0]
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    pic = scratch.shapes.add_picture(str(tmp_logo), prs.slide_width - Inches(2.6), Inches(0.5), Inches(2.4))
    blip = pic._element.xpath(".//a:blip")[0]
    image_part = scratch.part.related_part(blip.get(qn("r:embed")))
    blip.set(qn("r:embed"), master.part.relate_to(image_part, RT.IMAGE))
    master.shapes._spTree.append(pic._element)
    sld_id = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)
    tmp_logo.unlink()
    layout = prs.slide_layouts[1]
    for i in range(slides):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Slide {i + 1}: Project update"
        slide.placeholders[1].text = "Milestones reached\nBudget on track\nNext steps agreed"
        box = slide.shapes.add_textbox(Inches(0.2), Inches(0.05), Emu(prs.slide_width - Inches(0.4)), Inches(0.4))
        box.text_frame.text = "Evaluation Warning : The document was created with Spire.Presentation for Python"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    prs.save(path)


VML_WATERMARK = """
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
  <w:pPr><w:pStyle w:val="Header"/></w:pPr>
  <w:r><w:pict>
    <v:shape id="PowerPlusWaterMarkObject357922611" o:spid="_x0000_s2049" type="#_x0000_t136"
      style="position:absolute;margin-left:0;margin-top:0;width:412.4pt;height:247.45pt;rotation:315;z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
      o:allowincell="f" fillcolor="silver" stroked="f">
      <v:fill opacity=".5"/>
      <v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="DRAFT"/>
    </v:shape>
  </w:pict></w:r>
</w:p>
"""


def write_sample_docx(path: Path) -> None:
    from docx import Document
    from lxml import etree

    doc = Document()
    doc.add_heading("Service agreement", 1)
    for _ in range(4):
        doc.add_paragraph(BODY)
    header = doc.sections[0].header
    header.paragraphs[0].text = "Acme Corp."
    header._element.append(etree.fromstring(VML_WATERMARK))
    doc.save(path)


def main(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    _, tiled, _ = sample_tiled_image()
    cv2.imwrite(str(out_dir / "photo-tiled-watermark.jpg"), tiled, [cv2.IMWRITE_JPEG_QUALITY, 92])
    _, logo, _ = sample_logo_image()
    cv2.imwrite(str(out_dir / "photo-logo.jpg"), logo, [cv2.IMWRITE_JPEG_QUALITY, 92])
    write_sample_video(out_dir / "clip-moving-watermark.mp4", seconds=6)
    write_sample_pdf(out_dir / "report-confidential.pdf")
    write_sample_pptx(out_dir / "slides-evaluation.pptx")
    write_sample_docx(out_dir / "agreement-draft.docx")
    print("samples written to", out_dir)


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "samples")
