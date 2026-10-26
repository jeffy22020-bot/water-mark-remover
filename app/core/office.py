"""PowerPoint (.pptx) and Word (.docx) watermark removal.

Office watermarks are shapes: text boxes or pictures placed on the slide master/layouts
or repeated on every slide (PowerPoint), or "PowerPlusWaterMarkObject" shapes in page
headers (Word). We list them as candidates, the user picks, and we delete the shapes.

Pictures that have a watermark burned in can be opened in the image editor
(``get_media`` / ``replace_media``) and are swapped in place everywhere they are used.
"""

from __future__ import annotations

import base64
import hashlib
import io
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator, Optional

from PIL import Image

from .pdf import KEYWORDS

log = logging.getLogger(__name__)

PPT_EXTENSIONS = {".pptx", ".pptm", ".ppsx", ".potx"}
WORD_EXTENSIONS = {".docx", ".docm", ".dotx"}
LEGACY_EXTENSIONS = {".ppt", ".pps", ".pot", ".odp", ".doc", ".odt", ".rtf"}
EDITABLE_MEDIA = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "v": "urn:schemas-microsoft-com:vml",
    "o": "urn:schemas-microsoft-com:office:office",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
}
STRONG_WORDS = re.compile(r"water\s*mark|evaluation|created (with|by|using)|spire\.|aspose|trial version|unlicensed|unregistered", re.I)


@dataclass
class Candidate:
    id: str
    key: str
    kind: str  # text | picture | shape
    label: str
    where: set = field(default_factory=set)
    count: int = 0
    on_master: bool = False
    rotated: bool = False
    transparent: bool = False
    keyword: bool = False
    named: bool = False
    marker: bool = False
    thumb: Optional[str] = None
    score: float = 0.0
    reasons: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "id": self.id, "kind": self.kind, "label": self.label, "thumb": self.thumb,
            "where": sorted(self.where, key=_where_sort)[:12], "count": self.count,
            "score": round(self.score, 2), "reasons": self.reasons, "selected": self.score >= 0.6,
        }


def _where_sort(w: str):
    m = re.search(r"(\d+)", w)
    return (0 if w.startswith(("Master", "Layout", "Header", "Footer")) else 1, int(m.group(1)) if m else 0, w)


def _thumb(blob: bytes, size: int = 160) -> Optional[str]:
    try:
        with Image.open(io.BytesIO(blob)) as im:
            im.thumbnail((size, size))
            buf = io.BytesIO()
            im.convert("RGBA").save(buf, "PNG")
        return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    except Exception:  # noqa: BLE001 - EMF/WMF etc.
        return None


def kind_of(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext in PPT_EXTENSIONS:
        return "pptx"
    if ext in WORD_EXTENSIONS:
        return "docx"
    if ext in LEGACY_EXTENSIONS:
        return "legacy"
    raise ValueError("Unsupported document type")


# ------------------------------------------------------------------ PowerPoint
def _pptx_shapes(prs) -> Iterator[tuple[object, str, bool]]:
    """(shape, where, on_master) for every shape, recursing into groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes, where, on_master):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes, where, on_master)
            else:
                yield shape, where, on_master

    for mi, master in enumerate(prs.slide_masters, 1):
        yield from walk(master.shapes, f"Master {mi}", True)
        for layout in master.slide_layouts:
            yield from walk(layout.shapes, f"Layout “{layout.name}”", True)
    for si, slide in enumerate(prs.slides, 1):
        yield from walk(slide.shapes, f"Slide {si}", False)


def _pptx_shape_info(shape) -> Optional[dict]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.is_placeholder:
        return None
    el = shape._element
    xml = el.xml
    name = shape.name or ""
    text = shape.text_frame.text.strip() if shape.has_text_frame else ""
    rotated = (shape.rotation or 0) % 90 not in (0,)
    transparent = bool(re.search(r"<a:alpha val=\"(\d+)\"", xml)) or "alphaModFix" in xml
    info = {"name": name, "text": text, "rotated": rotated, "transparent": transparent, "thumb": None}
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            blob = shape.image.blob
        except Exception:  # noqa: BLE001 - linked pictures
            return None
        info.update(kind="picture", key="pic:" + hashlib.sha1(blob).hexdigest()[:16], blob=blob)
        info["label"] = f"Picture “{name}”" if name else "Picture"
    elif text:
        norm = re.sub(r"\s+", " ", text).strip().lower()[:200]
        info.update(kind="text", key="txt:" + hashlib.sha1(norm.encode()).hexdigest()[:16])
        info["label"] = "Text “" + re.sub(r"\s+", " ", text)[:70] + "”"
    elif re.search(r"water\s*mark", name, re.I) or (rotated and transparent):
        info.update(kind="shape", key="shp:" + hashlib.sha1(xml.encode()).hexdigest()[:16], label=f"Shape “{name}”")
    else:
        return None
    return info


def _analyze_pptx(path: str) -> dict:
    from pptx import Presentation

    prs = Presentation(path)
    slide_count = len(prs.slides)
    cands: dict[str, Candidate] = {}
    slides_with: dict[str, set] = {}
    for shape, where, on_master in _pptx_shapes(prs):
        info = _pptx_shape_info(shape)
        if info is None:
            continue
        c = cands.get(info["key"])
        if c is None:
            c = cands[info["key"]] = Candidate(f"o{len(cands) + 1}", info["key"], info["kind"], info["label"])
            if info.get("blob"):
                c.thumb = _thumb(info["blob"])
        c.where.add(where)
        c.count += 1
        c.on_master |= on_master
        c.rotated |= info["rotated"]
        c.transparent |= info["transparent"]
        c.keyword |= bool(KEYWORDS.search(info["text"])) or bool(STRONG_WORDS.search(info["text"]))
        c.marker |= bool(STRONG_WORDS.search(info["text"]))
        c.named |= bool(re.search(r"water\s*mark", info["name"], re.I))
        if not on_master:
            slides_with.setdefault(info["key"], set()).add(where)
    for key, c in cands.items():
        share = len(slides_with.get(key, ())) / max(1, slide_count)
        _score(c, share, slide_count, "slides")
    return _result("pptx", cands, _media_list(prs.part.package, "/ppt/media/"))


def _remove_pptx(path: str, out_path: str, keys: set[str]) -> None:
    from pptx import Presentation

    prs = Presentation(path)
    for shape, _, _ in list(_pptx_shapes(prs)):
        info = _pptx_shape_info(shape)
        if info and info["key"] in keys:
            el = shape._element
            el.getparent().remove(el)
    prs.save(out_path)


# ------------------------------------------------------------------ Word
def _docx_parts(doc) -> Iterator[tuple[object, str, bool]]:
    """(xml element, where, in_header) for headers, footers and the body."""
    seen = set()
    for si, section in enumerate(doc.sections, 1):
        for attr, label in (
            ("header", "Header"), ("first_page_header", "First-page header"), ("even_page_header", "Even-page header"),
            ("footer", "Footer"), ("first_page_footer", "First-page footer"), ("even_page_footer", "Even-page footer"),
        ):
            try:
                part = getattr(section, attr)
                if part.is_linked_to_previous:
                    continue
                el = part._element
            except Exception:  # noqa: BLE001
                continue
            if id(el) in seen:
                continue
            seen.add(id(el))
            yield el, f"{label} (section {si})", True
    yield doc.element.body, "Document body", False


def _docx_items(doc) -> Iterator[tuple[dict, object, str, bool]]:
    """Watermark-like elements: (info, element-to-remove, where, in_header)."""
    for root, where, in_header in _docx_parts(doc):
        # VML shapes (classic Word watermarks: text via textpath, pictures via imagedata).
        for pict in root.iter(f"{{{NS['w']}}}pict"):
            shape = next(pict.iter(f"{{{NS['v']}}}shape"), None)
            if shape is None:
                continue
            sid = shape.get("id", "") + " " + shape.get(f"{{{NS['o']}}}spid", "")
            textpath = next(shape.iter(f"{{{NS['v']}}}textpath"), None)
            text = textpath.get("string", "") if textpath is not None else ""
            named = "watermark" in sid.lower()
            if not (named or (in_header and (text or shape.find(f"{{{NS['v']}}}imagedata") is not None))):
                continue
            style = shape.get("style", "")
            info = {
                "kind": "text" if text else "picture",
                "label": f"Watermark “{text}”" if text else "Picture watermark",
                "key": "vml:" + hashlib.sha1((text or sid.split()[0] if sid.split() else text).encode()).hexdigest()[:16],
                "text": text,
                "named": named,
                "rotated": "rotation" in style,
                "transparent": "opacity" in pict.xml if hasattr(pict, "xml") else False,
            }
            yield info, _run_of(pict), where, in_header
        # DrawingML anchored shapes placed behind the text in headers.
        for anchor in root.iter(f"{{{NS['wp']}}}anchor"):
            doc_pr = anchor.find(f"{{{NS['wp']}}}docPr")
            name = doc_pr.get("name", "") if doc_pr is not None else ""
            behind = anchor.get("behindDoc") == "1"
            text = "".join(t.text or "" for t in anchor.iter(f"{{{NS['w']}}}t")).strip()
            named = "watermark" in name.lower()
            if not (named or (in_header and behind) or STRONG_WORDS.search(text)):
                continue
            info = {
                "kind": "text" if text else "picture",
                "label": (f"Shape “{text[:60]}”" if text else f"Picture “{name}”" if name else "Picture behind text"),
                "key": "dml:" + hashlib.sha1((name + "|" + text).encode()).hexdigest()[:16],
                "text": text, "named": named, "rotated": False, "transparent": behind,
            }
            yield info, _run_of(anchor), where, in_header
        # Evaluation banners inserted as plain paragraphs (Spire.Doc, Aspose.Words…).
        for para in root.iter(f"{{{NS['w']}}}p"):
            text = "".join(t.text or "" for t in para.iter(f"{{{NS['w']}}}t")).strip()
            if text and STRONG_WORDS.search(text) and len(text) < 300:
                info = {
                    "kind": "text", "label": "Text “" + text[:70] + "”",
                    "key": "par:" + hashlib.sha1(text.lower().encode()).hexdigest()[:16],
                    "text": text, "named": False, "rotated": False, "transparent": False,
                }
                yield info, para, where, in_header


def _run_of(el):
    """The enclosing w:r run (removing it removes the shape cleanly)."""
    node = el
    while node is not None and node.tag != f"{{{NS['w']}}}r":
        node = node.getparent()
    return node if node is not None else el


def _analyze_docx(path: str) -> dict:
    from docx import Document

    doc = Document(path)
    cands: dict[str, Candidate] = {}
    for info, _, where, in_header in _docx_items(doc):
        c = cands.get(info["key"])
        if c is None:
            c = cands[info["key"]] = Candidate(f"o{len(cands) + 1}", info["key"], info["kind"], info["label"])
        c.where.add(where)
        c.count += 1
        c.on_master |= in_header
        c.rotated |= info["rotated"]
        c.transparent |= info["transparent"]
        c.named |= info["named"]
        c.keyword |= bool(KEYWORDS.search(info["text"]))
        c.marker |= bool(STRONG_WORDS.search(info["text"]))
    for c in cands.values():
        _score(c, 0, 0, "pages")
    return _result("docx", cands, _media_list(doc.part.package, "/word/media/"))


def _remove_docx(path: str, out_path: str, keys: set[str]) -> None:
    from docx import Document

    doc = Document(path)
    doomed = [el for info, el, _, _ in _docx_items(doc) if info["key"] in keys]
    for el in doomed:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
    doc.save(out_path)


# ------------------------------------------------------------------ shared
def _score(c: Candidate, share: float, total: int, unit: str) -> None:
    score, reasons = 0.0, []
    if c.named:
        score += 1.0
        reasons.append("named “watermark”")
    if c.marker:
        score += 0.8
        reasons.append("watermark / evaluation wording")
    elif c.keyword:
        score += 0.5
        reasons.append("watermark wording")
    if c.on_master:
        score += 0.25
        reasons.append("on the master / header – shows on every " + ("slide" if unit == "slides" else "page"))
    if total >= 2 and share >= 0.5:
        score += 0.3
        reasons.append(f"repeated on {round(share * total)} of {total} {unit}")
    if c.rotated:
        score += 0.3
        reasons.append("rotated")
    if c.transparent:
        score += 0.3
        reasons.append("see-through" if unit == "slides" else "behind the text")
    c.score, c.reasons = score, reasons


def _result(doc_type: str, cands: dict, media: list) -> dict:
    listed = sorted((c for c in cands.values() if c.score >= 0.25), key=lambda c: -c.score)[:40]
    return {
        "type": doc_type,
        "candidates": [c.to_dict() for c in listed],
        "media": media,
        "_keys": {c.id: c.key for c in listed},
    }


def _media_list(package, prefix: str) -> list[dict]:
    items = []
    for part in package.iter_parts():
        name = str(part.partname)
        if not name.startswith(prefix):
            continue
        ext = Path(name).suffix.lower()
        if ext not in EDITABLE_MEDIA:
            continue
        blob = part.blob
        try:
            with Image.open(io.BytesIO(blob)) as im:
                width, height = im.size
        except Exception:  # noqa: BLE001
            continue
        if width < 48 or height < 48:
            continue
        items.append({"name": name, "width": width, "height": height, "thumb": _thumb(blob, 220)})
    return items


def _open_package(path: str):
    doc_type = kind_of(path)
    if doc_type == "pptx":
        from pptx import Presentation

        doc = Presentation(path)
    else:
        from docx import Document

        doc = Document(path)
    return doc


def analyze(path: str) -> dict:
    return _analyze_pptx(path) if kind_of(path) == "pptx" else _analyze_docx(path)


def remove(path: str, out_path: str, keys: set[str]) -> str:
    (_remove_pptx if kind_of(path) == "pptx" else _remove_docx)(path, out_path, keys)
    return out_path


def get_media(path: str, name: str) -> bytes:
    doc = _open_package(path)
    for part in doc.part.package.iter_parts():
        if str(part.partname) == name:
            return part.blob
    raise KeyError(name)


def replace_media(path: str, out_path: str, name: str, image_bytes: bytes) -> str:
    """Swap one embedded picture (re-encoded in the original format)."""
    doc = _open_package(path)
    for part in doc.part.package.iter_parts():
        if str(part.partname) == name:
            ext = Path(name).suffix.lower()
            fmt = {".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".gif": "GIF", ".bmp": "BMP",
                   ".tif": "TIFF", ".tiff": "TIFF", ".webp": "WEBP"}[ext]
            with Image.open(io.BytesIO(image_bytes)) as im:
                with Image.open(io.BytesIO(part.blob)) as original:
                    keep_alpha = original.mode in ("RGBA", "LA", "P") and fmt in ("PNG", "WEBP", "TIFF")
                im = im.convert("RGBA" if keep_alpha else "RGB")
                buf = io.BytesIO()
                im.save(buf, fmt, **({"quality": 95} if fmt == "JPEG" else {}))
            part._blob = buf.getvalue()
            doc.save(out_path)
            return out_path
    raise KeyError(name)


# ------------------------------------------------------------------ legacy formats
def find_libreoffice() -> Optional[str]:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    if sys.platform == "win32":
        for base in (os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            candidate = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if candidate.exists():
                return str(candidate)
    elif sys.platform == "darwin":
        candidate = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if candidate.exists():
            return str(candidate)
    return None


def convert_legacy(path: str, out_dir: str) -> str:
    """Convert .ppt/.doc/.odp/... to .pptx/.docx with LibreOffice (if installed)."""
    soffice = find_libreoffice()
    ext = Path(path).suffix.lower()
    target = "pptx" if ext in {".ppt", ".pps", ".pot", ".odp"} else "docx"
    if not soffice:
        raise RuntimeError(
            f"Old “{ext}” files must be converted first. Open the file in PowerPoint/Word and choose "
            f"Save As → .{target}, or install the free LibreOffice and try again."
        )
    from .ffmpeg import popen_kwargs

    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [soffice, f"-env:UserInstallation=file:///{Path(profile).as_posix().lstrip('/')}", "--headless",
             "--convert-to", target, "--outdir", out_dir, path],
            stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300,
            **popen_kwargs(),
        )
    out = Path(out_dir) / (Path(path).stem + "." + target)
    if not out.exists():
        raise RuntimeError("LibreOffice could not convert this file.")
    return str(out)
